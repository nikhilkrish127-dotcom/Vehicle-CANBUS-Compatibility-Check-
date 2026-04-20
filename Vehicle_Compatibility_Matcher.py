
from pathlib import Path
import re
from collections import defaultdict

import pandas as pd
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

BASE = Path(__file__).resolve().parent
WORKBOOK = BASE / "Vehicle_Compatibility_Workbench_FINAL_v9_clean_report.xlsx"
FINAL_REPORT_XLSX = BASE / "Final_Report_Clean.xlsx"
FINAL_REPORT_CSV = BASE / "Final_Report_Clean_Supported.csv"

GEN_RE = re.compile(r"\b(mp[1-6]|mk[1-6]|j[123]\d{2}|w\d{3}|g\d{2,3}|xzu[a-z0-9]+|axzh\d+|axza\d+|gsz\d+|vs30|w447|tq|tga|tgs|tgx|gmt1xx|spa|tnga-f|clar|ld)\b", re.I)

STOP_TOKENS = {
    "type","facelift","electric","hybrid","diesel","petrol","phev","ev","ice",
    "wagon","sedan","truck","trucks","van","bus","buses","rigid","tractor","head",
    "4x2","4x4","6x4","6x2","7.5","ton","sport","tourer","signature","limited",
    "luxury","premium","similar","or","and","car","cars","coach","minibus",
    "pickup","pick","up","suv","hatchback","saloon","vehicle","asset","description",
    "passenger","staff","school","cargo"
}
CATEGORY_WORDS = {
    "bus","buses","coach","minibus","truck","trucks","tractor","rigid","tipper","hauler","lorry",
    "van","minivan","panel","car","cars","sedan","saloon","wagon","suv","hatchback","coupe","convertible","estate",
    "pickup","pick","up","cab","doublecab","double","crewcab","crew"
}
GENERIC_FAMILY_TERMS = {
    "", "bus", "truck", "van", "car", "pickup", "coach", "minibus", "staff bus", "school bus"
}

def norm_text(s):
    if pd.isna(s):
        return ""
    s = str(s).strip().lower()
    s = s.replace("&", " and ")
    s = re.sub(r"[\(\)\[\]\{\}/\\,._\-]+", " ", s)
    s = re.sub(r"[^a-z0-9+\s]", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

def norm_make(s):
    s = norm_text(s)
    repl = {
        "mercedes benz":"mercedesbenz","mercedes":"mercedesbenz","benz":"mercedesbenz",
        "nissan ud":"ud","ud trucks":"ud","land rover":"landrover","great wall":"greatwall",
        "alfa romeo":"alfaromeo","rolls royce":"rollsroyce","vw":"volkswagen","chevy":"chevrolet",
        "ashok leyland":"ashokleyland","ashoka leyland":"ashokleyland",
        "bharat benz":"bharatbenz","force motors":"forcemotors","general motors":"gm",
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    return s.replace(" ", "")

def extract_year(s):
    if pd.isna(s):
        return None
    m = re.search(r"(19|20)\d{2}", str(s))
    return int(m.group(0)) if m else None

def extract_generation(text):
    txt = norm_text(text)
    hits = GEN_RE.findall(txt)
    return ", ".join(sorted(set([h.lower() for h in hits]))) if hits else ""

def parse_year_ranges(year_text):
    if pd.isna(year_text):
        return []
    s = str(year_text).replace("–","-").replace("—","-").replace(" ","")
    s = s.replace(">=", "+").replace(">", "+")
    ranges = []
    for part in re.split(r"[;,/]", s):
        if not part:
            continue
        m = re.fullmatch(r"(\d{4})\+", part)
        if m:
            ranges.append((int(m.group(1)), None)); continue
        m = re.fullmatch(r"(\d{4})-(\d{4})", part)
        if m:
            ranges.append((int(m.group(1)), int(m.group(2)))); continue
        m = re.fullmatch(r"(\d{4})", part)
        if m:
            y = int(m.group(1)); ranges.append((y, y)); continue
        m = re.search(r"(\d{4})-(\d{4})", part)
        if m:
            ranges.append((int(m.group(1)), int(m.group(2)))); continue
        m = re.search(r"(\d{4})\+", part)
        if m:
            ranges.append((int(m.group(1)), None)); continue
    return ranges

def year_in_range(year, year_text):
    if year is None:
        return None
    ranges = parse_year_ranges(year_text)
    if not ranges:
        return None
    for start, end in ranges:
        if end is None and year >= start:
            return True
        if end is not None and start <= year <= end:
            return True
    return False

def year_distance(year, year_text):
    if year is None:
        return None
    ranges = parse_year_ranges(year_text)
    if not ranges:
        return None
    distances = []
    for start, end in ranges:
        if end is None:
            distances.append(0 if year >= start else start - year)
        else:
            if start <= year <= end:
                distances.append(0)
            elif year < start:
                distances.append(start - year)
            else:
                distances.append(year - end)
    return min(distances) if distances else None

def normalize_category(val):
    txt = norm_text(val)
    if not txt:
        return ""
    if any(t in txt for t in ["bus","buses","coach","minibus"]):
        return "bus"
    if any(t in txt for t in ["truck","trucks","tractor","rigid","tipper","hauler","lorry"]):
        return "truck"
    if any(t in txt for t in ["van","minivan","panel"]):
        return "van"
    if any(t in txt for t in ["pickup","double cab","doublecab","crew cab","crewcab"]):
        return "pickup"
    if any(t in txt for t in ["car","cars","sedan","saloon","wagon","suv","hatchback","coupe","convertible","estate"]):
        return "car"
    if any(t in txt for t in ["motorbike","motorcycle"]):
        return "motorbike"
    if any(t in txt for t in ["agricultural"]):
        return "agri"
    if any(t in txt for t in ["machinery","construction","utility","forest","heavy"]):
        return "machinery"
    if any(t in txt for t in ["motorboats","jet ski","snowmobile","marine","boat"]):
        return "marine"
    return ""

def infer_input_category(raw_desc):
    return normalize_category(raw_desc)

def infer_model_family(model_text):
    txt = norm_text(model_text)
    if not txt:
        return ""
    toks = [t for t in re.split(r"\s+", txt) if t]
    kept = []
    for i, t in enumerate(toks):
        if re.fullmatch(r"(19|20)\d{2}\+?", t):
            continue
        if t in STOP_TOKENS:
            continue
        if re.fullmatch(r"mp[1-6]|mk[1-6]|j[123]\d{2}|w\d{3}|g\d{2,3}|vs30|w447|tq|tga|tgs|tgx|gmt1xx|spa|clar|ld", t):
            continue
        if re.fullmatch(r"\d+", t):
            if i == 0:
                kept.append(t)
            continue
        kept.append(t)
        if len(kept) >= 3:
            break
    return " ".join(kept[:3]).strip()

def detect_fuel_support(param_text):
    if pd.isna(param_text):
        return "No"
    text = str(param_text).lower()
    fuel_keywords = [
        "fuel level","fuel level (%)","fuel level 1 (l)","fuel level 2 (l)","fuel level (l)",
        "fuel (l)","fuel volume","fuel consumption","instant fuel consumption",
        "average fuel consumption","fuel used","fuel rate","fuel flow","fuel cc",
    ]
    return "Yes" if any(k in text for k in fuel_keywords) else "No"

def reorder_parameters_for_display(param_text):
    if pd.isna(param_text):
        return ""
    parts = [p.strip() for p in str(param_text).split(";") if p.strip()]
    if not parts:
        return ""
    fuel_first, others = [], []
    for p in parts:
        (fuel_first if "fuel" in p.lower() else others).append(p)
    return "; ".join(fuel_first + others)

def autosize_sheet(ws):
    for col_cells in ws.columns:
        values = ["" if c.value is None else str(c.value) for c in col_cells]
        max_len = min(60, max((len(v) for v in values), default=0) + 2)
        ws.column_dimensions[col_cells[0].column_letter].width = max(10, max_len)

def prepare_report_sheet(wb, name, headers):
    if name in wb.sheetnames:
        ws = wb[name]
        wb.remove(ws)
    ws = wb.create_sheet(name)
    for idx, h in enumerate(headers, start=1):
        ws.cell(1, idx).value = h
    return ws

def update_table_ref(ws, table_name):
    from openpyxl.utils import get_column_letter
    if table_name in ws.tables:
        ws.tables[table_name].ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

def color_results(ws, status_col_letter):
    fills = {
        "Strong Match": PatternFill("solid", fgColor="E2F0D9"),
        "Possible Match": PatternFill("solid", fgColor="FFF2CC"),
        "Review Needed": PatternFill("solid", fgColor="FCE4D6"),
        "No Reliable Match": PatternFill("solid", fgColor="F4CCCC"),
    }
    for r in range(2, ws.max_row + 1):
        val = ws[f"{status_col_letter}{r}"].value
        if val in fills:
            for c in ws[r]:
                c.fill = fills[val]

def normalize_master(master):
    master = master.copy()
    needed = ["Brand","Model","Model_Family","Model_Year_Text","Support_List_Type","Source_Sheet",
              "Brand_Norm","Model_Norm","Family_Norm","Generation_Hint","Supported_Parameter_Count",
              "Supported_Parameters_Sample","Vehicle_Category"]
    for c in needed:
        if c not in master.columns:
            master[c] = ""
    master["Brand_Norm"] = master["Brand_Norm"].fillna("").replace("nan","")
    master.loc[master["Brand_Norm"]=="","Brand_Norm"] = master["Brand"].map(norm_make)
    master["Model_Norm"] = master["Model_Norm"].fillna("").replace("nan","")
    master.loc[master["Model_Norm"]=="","Model_Norm"] = master["Model"].map(norm_text)
    master["Family_Norm"] = master["Family_Norm"].fillna("").replace("nan","")
    master.loc[master["Family_Norm"]=="","Family_Norm"] = master["Model_Family"].map(norm_text)
    master["Generation_Hint"] = master["Generation_Hint"].fillna("").replace("nan","")
    master.loc[master["Generation_Hint"]=="","Generation_Hint"] = master["Model"].map(extract_generation)
    master["Category_Bucket"] = master["Vehicle_Category"].apply(normalize_category)
    master["Full_Supported_Parameters"] = master["Supported_Parameters_Sample"].fillna("").astype(str).apply(reorder_parameters_for_display)
    master["Param_Preview"] = master["Full_Supported_Parameters"].fillna("").astype(str).apply(lambda x: "; ".join([p.strip() for p in x.split(";") if p.strip()][:5]))
    master["Fuel_Data_Available"] = master["Supported_Parameters_Sample"].apply(detect_fuel_support)
    return master

def load_sheet_df(sheet_name):
    return pd.read_excel(WORKBOOK, sheet_name=sheet_name)

def build_alias_maps(alias_df):
    make_aliases = []
    model_aliases = []
    category_aliases = []
    if alias_df is None or alias_df.empty:
        return make_aliases, model_aliases, category_aliases
    adf = alias_df.copy()
    adf["Rule_Type"] = adf["Rule_Type"].astype(str)
    if "Priority" in adf.columns:
        adf["Priority"] = pd.to_numeric(adf["Priority"], errors="coerce").fillna(0)
    else:
        adf["Priority"] = 0
    adf = adf.sort_values(["Priority"], ascending=False)
    for _, r in adf.iterrows():
        rec = {
            "src": norm_text(r.get("Input_Text","")),
            "dst": norm_text(r.get("Normalized_Output","")),
            "priority": r.get("Priority",0),
            "notes": str(r.get("Notes","") or "")
        }
        rt = str(r.get("Rule_Type","")).strip()
        if rt == "Make_Alias":
            make_aliases.append(rec)
        elif rt == "Model_Alias":
            model_aliases.append(rec)
        elif rt == "Category_Alias":
            category_aliases.append(rec)
    return make_aliases, model_aliases, category_aliases

def detect_make(raw_norm, master, make_aliases):
    joined = raw_norm.replace(" ", "")
    joined_make = norm_make(raw_norm)
    for r in make_aliases:
        src = r["src"]
        dst = r["dst"]
        if src and src in raw_norm:
            return norm_make(dst)
    brand_map = master[["Brand","Brand_Norm"]].dropna().drop_duplicates()
    candidates = []
    for _, br in brand_map.iterrows():
        brand_label = norm_text(br["Brand"])
        brand_norm = str(br["Brand_Norm"])
        if not brand_norm:
            continue
        if brand_label and re.search(rf"\b{re.escape(brand_label)}\b", raw_norm):
            return brand_norm
        if brand_norm in joined or brand_norm in joined_make:
            return brand_norm
        candidates.append((brand_label, brand_norm))
    # fuzzy brand rescue for typos like Ashoka Leyland
    prefix_tokens = []
    for tok in raw_norm.split():
        if re.fullmatch(r"(19|20)\d{2}", tok):
            continue
        if tok in CATEGORY_WORDS:
            continue
        prefix_tokens.append(tok)
        if len(prefix_tokens) >= 3:
            break
    prefix = " ".join(prefix_tokens)
    prefix_norm = norm_make(prefix)
    if not prefix:
        return ""
    scored = []
    for brand_label, brand_norm in candidates:
        score = max(fuzz.ratio(prefix, brand_label), fuzz.ratio(prefix_norm, brand_norm))
        if brand_label and prefix.startswith(brand_label):
            score += 8
        scored.append((score, brand_norm))
    scored.sort(reverse=True)
    if scored and scored[0][0] >= 90 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 4):
        return scored[0][1]
    return ""

def remove_detected_brand(raw_desc, parsed_make, master):
    txt = norm_text(raw_desc)
    if not parsed_make:
        return txt
    brand_variants = master.loc[master["Brand_Norm"] == parsed_make, "Brand"].dropna().astype(str).map(norm_text).unique().tolist()
    brand_variants = [b for b in brand_variants if b]
    cleaned = txt
    for bv in sorted(set(brand_variants), key=len, reverse=True):
        cleaned = re.sub(rf"\b{re.escape(bv)}\b", " ", cleaned)

    # fuzzy prefix trim for variants/typos like Ashoka Leyland vs Ashok Leyland
    raw_tokens = txt.split()
    best_n = 0
    best_score = 0
    for bv in brand_variants:
        bv_tokens = bv.split()
        low = max(1, len(bv_tokens) - 1)
        high = min(len(raw_tokens), len(bv_tokens) + 1)
        for n in range(low, high + 1):
            prefix = " ".join(raw_tokens[:n])
            score = max(fuzz.token_set_ratio(prefix, bv), fuzz.ratio(norm_make(prefix), parsed_make))
            if score > best_score:
                best_score = score
                best_n = n
    if best_score >= 88 and best_n > 0:
        cleaned = " ".join(raw_tokens[best_n:])

    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned

def build_search_context(raw_desc, parsed_make, input_category, category_aliases=None):
    txt = norm_text(raw_desc)
    txt = re.sub(r"(19|20)\d{2}", " ", txt)
    if parsed_make:
        txt = re.sub(re.escape(parsed_make), " ", txt)
    for cat in ["bus","buses","coach","minibus","truck","trucks","tractor","rigid","tipper","lorry",
                "van","minivan","panel","car","cars","sedan","saloon","wagon","suv","hatchback","pickup","cab"]:
        txt = re.sub(rf"\b{re.escape(cat)}\b", " ", txt)
    if category_aliases:
        for r in category_aliases:
            if r["src"] and r["src"] in txt:
                txt = txt.replace(r["src"], " ")
    txt = re.sub(r"\s+", " ", txt).strip()
    return txt

def special_family_rules(raw_desc, parsed_make):
    txt = norm_text(raw_desc)
    if parsed_make == "toyota":
        if "land cruiser" in txt: return "land cruiser"
        if "corolla cross" in txt: return "corolla cross"
        if "corolla" in txt: return "corolla"
        if "yaris" in txt: return "yaris"
        if "hilux" in txt: return "hilux"
        if "coaster" in txt: return "coaster"
    if parsed_make == "mercedesbenz":
        if "actros" in txt: return "actros"
        if "sprinter" in txt: return "sprinter"
        if "vito" in txt: return "vito"
        if "citaro" in txt: return "citaro"
    if parsed_make == "ashokleyland":
        if "dost" in txt: return "dost"
        if "partner" in txt: return "partner"
        if "ecomet" in txt: return "ecomet"
        if "boss" in txt: return "boss"
        if "viking" in txt: return "viking"
        if "lynx" in txt: return "lynx"
        if "falcon" in txt: return "falcon"
        if "cheetah" in txt: return "cheetah"
        if "oyster" in txt: return "oyster"
    if parsed_make == "volvo":
        if "xc90" in txt: return "xc90"
        if "s90" in txt: return "s90"
        if re.search(r"\bfl\b", txt): return "fl"
    if parsed_make == "man":
        if "tgs" in txt: return "tgs"
        if "tga" in txt: return "tga"
        if "tgx" in txt: return "tgx"
    if parsed_make == "hino" and re.search(r"\bxzu", txt): return "xzu"
    return ""

def extract_parsed_model_family(raw_desc, parsed_make, master, model_aliases):
    special = special_family_rules(raw_desc, parsed_make)
    if special:
        return special
    raw_norm = norm_text(raw_desc)
    for r in model_aliases:
        src = r["src"]; dst = r["dst"]
        notes = r["notes"].lower()
        brand_ok = (not parsed_make) or (not notes) or (parsed_make in norm_make(notes)) or (parsed_make in notes.replace(" ",""))
        if src and src in raw_norm and dst and brand_ok:
            return dst
    temp = remove_detected_brand(raw_desc, parsed_make, master)
    temp = re.sub(r"(19|20)\d{2}", " ", temp)
    temp = re.sub(r"\s+", " ", temp).strip()
    if not temp:
        return ""
    fams = master.loc[master["Brand_Norm"] == parsed_make, "Family_Norm"].dropna().astype(str).unique().tolist() if parsed_make else master["Family_Norm"].dropna().astype(str).unique().tolist()
    fams = [f for f in fams if f and f != "nan"]
    best, best_score = "", -1
    for fam in fams:
        score = fuzz.token_set_ratio(temp, fam)
        if fam and fam in temp:
            score += 20
        if score > best_score:
            best, best_score = fam, score
    if best_score >= 72:
        return best
    inferred = infer_model_family(temp)
    if inferred and norm_make(inferred) == parsed_make:
        return ""
    return inferred

def check_override(raw_desc, overrides):
    raw_norm = norm_text(raw_desc)
    if "Active" not in overrides.columns:
        return None
    active = overrides[overrides["Active"].astype(str).str.strip().str.lower() == "yes"].copy()
    for _, r in active.iterrows():
        needle = norm_text(r.get("Raw_Match_Text", ""))
        if needle and needle in raw_norm:
            return r.to_dict()
    return None

def informative_tokens(text):
    toks = []
    for t in norm_text(text).split():
        if t in STOP_TOKENS or t in CATEGORY_WORDS:
            continue
        if re.fullmatch(r"(19|20)\d{2}", t):
            continue
        if len(t) <= 1:
            continue
        toks.append(t)
    return toks

def compute_model_similarity(search_text, parsed_family, row):
    row_family = str(row["Family_Norm"])
    row_model = str(row["Model_Norm"])
    info_tokens = set(informative_tokens(search_text or parsed_family))
    cand_tokens = set(informative_tokens(f"{row_family} {row_model}"))
    overlap = len(info_tokens & cand_tokens)
    coverage = overlap / max(1, len(info_tokens)) if info_tokens else 0
    fam_ratio = fuzz.token_set_ratio(parsed_family, row_family) if parsed_family else 0
    model_ratio = fuzz.token_set_ratio(search_text, row_model) if search_text else 0
    family_to_model = fuzz.token_set_ratio(search_text, row_family) if search_text else 0
    best_ratio = max(fam_ratio, model_ratio, family_to_model)

    score = 0
    model_check = "Weak"
    if parsed_family and parsed_family == row_family:
        score += 34; model_check = "Exact"
    elif parsed_family and (parsed_family in row_family or row_family in parsed_family):
        score += 24; model_check = "Close"
    elif coverage >= 0.80 and best_ratio >= 80:
        score += 22; model_check = "Close"
    elif coverage >= 0.55 and best_ratio >= 72:
        score += 14; model_check = "Family"
    elif best_ratio >= 78:
        score += 12; model_check = "Family"
    elif best_ratio >= 66:
        score += 4; model_check = "Weak"
    else:
        score -= 10; model_check = "Weak"

    if row_family and row_family in (search_text or ""):
        score += 8
    elif row_model and row_model in (search_text or ""):
        score += 6
    return round(score, 2), model_check, round(best_ratio, 1), round(coverage, 2)

def generic_input_flag(parsed_family, search_text):
    pf = norm_text(parsed_family)
    st = norm_text(search_text)
    info = informative_tokens(st)
    if pf in GENERIC_FAMILY_TERMS:
        return True
    if not pf and len(info) <= 1:
        return True
    if pf and pf in {"bus","truck","van","car","pickup"}:
        return True
    return False

def score_candidate(search_text, parsed_make, parsed_family, parsed_year, parsed_gen, input_category, row):
    row_brand_norm = str(row["Brand_Norm"])
    if not parsed_make or parsed_make != row_brand_norm:
        return None

    brand_check = "Exact"
    brand_points = 45

    row_cat = str(row.get("Category_Bucket","") or "")
    if input_category:
        if row_cat == input_category:
            category_check = "Exact"
            category_points = 14
        elif row_cat == "":
            category_check = "Unknown"
            category_points = 2
        else:
            category_check = "Mismatch"
            category_points = -18
    else:
        category_check = "Not provided"
        category_points = 3 if row_cat else 0

    model_points, model_check, best_ratio, coverage = compute_model_similarity(search_text, parsed_family, row)

    gen_points = 0
    row_gen = norm_text(row.get("Generation_Hint",""))
    if parsed_gen:
        if row_gen and any(code.strip() in row_gen for code in parsed_gen.split(",")):
            gen_points = 8
        elif row_gen:
            gen_points = -4

    ychk = year_in_range(parsed_year, row.get("Model_Year_Text", ""))
    ydist = year_distance(parsed_year, row.get("Model_Year_Text", ""))
    if parsed_year is None:
        year_check = "Not provided"
        year_points = 4
    elif ychk is True:
        year_check = "In range"
        year_points = 24
    elif ychk is False:
        if ydist is None:
            year_check = "Out of range"
            year_points = -18
        else:
            year_check = f"Out of range ({ydist}y)"
            year_points = -10 if ydist == 1 else (-18 if ydist <= 3 else -28)
    else:
        year_check = "No list year"
        year_points = 1

    score = round(brand_points + category_points + model_points + gen_points + year_points, 2)
    return {
        "score": score,
        "brand_check": brand_check,
        "model_check": model_check,
        "year_check": year_check,
        "category_check": category_check,
        "best_ratio": best_ratio,
        "coverage": coverage,
    }

def classify_candidate(cand, rows_for_brand, parsed_family, search_text):
    status = "Review Needed"
    confidence = "Low"
    if not cand:
        return "No Reliable Match", "Low"
    if cand["category_check"] == "Mismatch":
        return "No Reliable Match", "Low"
    generic = generic_input_flag(parsed_family, search_text)
    distinct_families = rows_for_brand["Family_Norm"].replace("", pd.NA).dropna().nunique() if isinstance(rows_for_brand, pd.DataFrame) and not rows_for_brand.empty else 0
    if cand["score"] >= 86 and cand["model_check"] in {"Exact","Close"} and cand["year_check"] in {"In range","Not provided","No list year"}:
        status, confidence = "Strong Match", "High"
    elif cand["score"] >= 68 and cand["model_check"] in {"Exact","Close","Family"}:
        status, confidence = "Possible Match", "Medium"
    elif cand["score"] >= 58 and cand["model_check"] != "Weak" and cand["category_check"] != "Mismatch":
        status, confidence = "Review Needed", "Medium"
    else:
        status, confidence = "No Reliable Match", "Low"
    if generic and distinct_families > 1 and status in {"Possible Match","Review Needed"} and cand["score"] < 82:
        status, confidence = "Review Needed", "Low"
    return status, confidence

def overall_reason(brand_check, model_check, year_check, category_check):
    return f"Brand: {brand_check}; Category: {category_check}; Model: {model_check}; Year: {year_check}"

def action_hint(status, parsed_year):
    if status == "Strong Match":
        return "Auto-accept candidate"
    if status == "Possible Match":
        return "Review top shortlist before approving"
    if status == "Review Needed" and parsed_year is None:
        return "Ask for exact model or year"
    if status == "Review Needed":
        return "Manual review recommended"
    return "No reliable match - collect more specific model details"

def prepare_candidate_pool(rows, input_category, search_text):
    if rows.empty:
        return rows
    # quick prefilter for performance and precision
    if input_category:
        same_cat = rows[rows["Category_Bucket"] == input_category].copy()
        if not same_cat.empty:
            rows = same_cat
    rows = rows.copy()
    rows["_prefilter_ratio"] = rows["Family_Norm"].fillna("").astype(str).apply(lambda x: fuzz.token_set_ratio(search_text, x) if search_text else 0)
    rows["_prefilter_ratio2"] = rows["Model_Norm"].fillna("").astype(str).apply(lambda x: fuzz.token_set_ratio(search_text, x) if search_text else 0)
    rows["_prefilter"] = rows[["_prefilter_ratio","_prefilter_ratio2"]].max(axis=1)
    if len(rows) > 250:
        rows = rows.sort_values(["_prefilter"], ascending=False).head(250).copy()
    return rows

def build_business_report(records, input_rows):
    def rec_action(status, confidence, supported, fuel_support, parsed_model):
        notes = []
        if supported == "No":
            notes.append("Check against support lists manually or request clearer model details")
        if status in ["Review Needed", "No Reliable Match"]:
            notes.append("Validate model wording before relying on compatibility")
        if fuel_support == "No" and supported == "Yes":
            notes.append("Supported vehicle but fuel level is unavailable on matched lists")
        if not parsed_model or str(parsed_model).strip() in GENERIC_FAMILY_TERMS:
            notes.append("Model description is generic; manual review recommended")
        if confidence == "Low" and supported == "Yes":
            notes.append("Low-confidence support result; confirm against source lists")
        return " | ".join(dict.fromkeys(notes)) if notes else "No immediate review needed"

    if not records:
        metrics = {
            "Total vehicles in list": 0,
            "Total brands in list": 0,
            "Total unique models": 0,
            "Vehicles supported by one or more lists but no fuel reading": 0,
            "Vehicles not included in any support list": 0,
            "Support coverage %": 0,
            "Strong match vehicles": 0,
            "Possible match vehicles": 0,
            "Review-needed vehicles": 0,
            "Duplicate input rows": 0,
        }
        empty_final = pd.DataFrame(columns=[
            "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Input_Category",
            "Supported","Best_Status","Confidence","Compatible_Adaptors","Any_Fuel_Support",
            "Matched_Candidates","Supported_Year_Ranges","Recommended_Review_Action"
        ])
        empty_unsupported = pd.DataFrame(columns=[
            "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Input_Category",
            "Best_Status","Confidence","Reason_Summary","Recommended_Review_Action"
        ])
        empty_no_fuel = pd.DataFrame(columns=[
            "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Compatible_Adaptors",
            "Matched_Candidates","Reason_Summary","Recommended_Review_Action"
        ])
        empty_brand_gap = pd.DataFrame(columns=["Brand","Vehicles_Not_In_Support_List","Unique_Models_Not_In_Support_List","Review_Recommendation"])
        empty_review = pd.DataFrame(columns=[
            "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Input_Category",
            "Best_Status","Confidence","Matched_Candidates","Reason_Summary","Recommended_Review_Action"
        ])
        return metrics, empty_final, empty_unsupported, empty_no_fuel, empty_brand_gap, empty_review

    df = pd.DataFrame(records)
    input_counts = pd.Series(input_rows).value_counts()

    grouped_rows = []
    for vehicle, g in df.groupby("Raw_Vehicle_Description", sort=True):
        input_count = int(input_counts.get(vehicle, len(g)))
        g_sorted = g.sort_values(["Match_Score"], ascending=False).copy()
        best_any = g_sorted.iloc[0]
        supported = g_sorted[g_sorted["Status"].isin(["Strong Match", "Possible Match"])]
        if supported.empty:
            row = {
                "Vehicle_Input": vehicle,
                "Input_Count": input_count,
                "Parsed_Brand": best_any.get("Parsed_Brand",""),
                "Parsed_Model_Family": best_any.get("Parsed_Model_Family",""),
                "Parsed_Year": best_any.get("Parsed_Year",""),
                "Input_Category": best_any.get("Input_Category",""),
                "Supported": "No",
                "Best_Status": best_any.get("Status",""),
                "Confidence": best_any.get("Confidence",""),
                "Compatible_Adaptors": "",
                "Any_Fuel_Support": "No",
                "Matched_Candidates": f'{best_any.get("Matched_Brand","")} {best_any.get("Matched_Model","")}'.strip(),
                "Supported_Year_Ranges": str(best_any.get("Supported_Year_Range","")),
                "Reason_Summary": best_any.get("Reason",""),
            }
        else:
            best_supported = supported.iloc[0]
            adaptors = sorted(set(supported["Source"].astype(str)))
            fuel_adaptors = sorted(set(supported.loc[supported["Fuel_Data_Available"].astype(str).str.upper()=="YES", "Source"].astype(str)))
            matched_candidates = sorted(set((supported["Matched_Brand"].astype(str) + " " + supported["Matched_Model"].astype(str)).str.strip()))
            year_ranges = sorted(set(supported["Supported_Year_Range"].astype(str)))
            row = {
                "Vehicle_Input": vehicle,
                "Input_Count": input_count,
                "Parsed_Brand": best_supported.get("Parsed_Brand",""),
                "Parsed_Model_Family": best_supported.get("Parsed_Model_Family",""),
                "Parsed_Year": best_supported.get("Parsed_Year",""),
                "Input_Category": best_supported.get("Input_Category",""),
                "Supported": "Yes",
                "Best_Status": best_supported.get("Status",""),
                "Confidence": best_supported.get("Confidence",""),
                "Compatible_Adaptors": ", ".join(adaptors),
                "Any_Fuel_Support": "Yes" if fuel_adaptors else "No",
                "Matched_Candidates": " | ".join(matched_candidates[:10]),
                "Supported_Year_Ranges": " | ".join([x for x in year_ranges if x]),
                "Reason_Summary": best_supported.get("Reason",""),
            }
        row["Recommended_Review_Action"] = rec_action(row["Best_Status"], row["Confidence"], row["Supported"], row["Any_Fuel_Support"], row["Parsed_Model_Family"])
        grouped_rows.append(row)

    final_report = pd.DataFrame(grouped_rows).sort_values(["Supported","Vehicle_Input"], ascending=[True, True]).reset_index(drop=True)
    # reorder with unsupported first
    final_report["Supported_Sort"] = final_report["Supported"].map({"No":0,"Yes":1}).fillna(2)
    final_report = final_report.sort_values(["Supported_Sort","Parsed_Brand","Vehicle_Input"]).drop(columns=["Supported_Sort"]).reset_index(drop=True)

    unsupported_df = final_report[final_report["Supported"] == "No"][ [
        "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Input_Category",
        "Best_Status","Confidence","Reason_Summary","Recommended_Review_Action"
    ] ].copy()

    no_fuel_df = final_report[(final_report["Supported"] == "Yes") & (final_report["Any_Fuel_Support"] == "No")][ [
        "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Compatible_Adaptors",
        "Matched_Candidates","Reason_Summary","Recommended_Review_Action"
    ] ].copy()

    review_df = final_report[(final_report["Supported"] == "No") | (final_report["Best_Status"].isin(["Review Needed", "No Reliable Match"])) | (final_report["Confidence"] == "Low") | (final_report["Any_Fuel_Support"] == "No") | (final_report["Parsed_Model_Family"].astype(str).isin(GENERIC_FAMILY_TERMS))][ [
        "Vehicle_Input","Input_Count","Parsed_Brand","Parsed_Model_Family","Parsed_Year","Input_Category",
        "Best_Status","Confidence","Matched_Candidates","Reason_Summary","Recommended_Review_Action"
    ] ].drop_duplicates().copy()

    brand_gap = unsupported_df.copy()
    if brand_gap.empty:
        brand_gap_df = pd.DataFrame(columns=["Brand","Vehicles_Not_In_Support_List","Unique_Models_Not_In_Support_List","Review_Recommendation"])
    else:
        brand_gap["Parsed_Brand"] = brand_gap["Parsed_Brand"].replace("", "Unrecognized")
        brand_gap["Model_For_Group"] = brand_gap["Parsed_Model_Family"].fillna("").replace("", "Unclear Model")
        brand_gap_df = brand_gap.groupby("Parsed_Brand", dropna=False).agg(
            Vehicles_Not_In_Support_List=("Vehicle_Input", "nunique"),
            Unique_Models_Not_In_Support_List=("Model_For_Group", lambda s: ", ".join(sorted(dict.fromkeys([x for x in s if x]))[:25]))
        ).reset_index().rename(columns={"Parsed_Brand":"Brand"})
        brand_gap_df["Review_Recommendation"] = brand_gap_df.apply(lambda r: "Expand support-list coverage or add mapping rules for this brand" if r["Brand"] != "Unrecognized" else "Review brand detection and raw input quality", axis=1)
        brand_gap_df = brand_gap_df.sort_values(["Vehicles_Not_In_Support_List","Brand"], ascending=[False, True])

    total_input_rows = int(len(input_rows))
    unique_entries = int(pd.Series(input_rows).nunique()) if input_rows else 0
    unique_brands = int(final_report["Parsed_Brand"].replace("", pd.NA).dropna().nunique()) if not final_report.empty else 0
    unique_models = int(final_report["Parsed_Model_Family"].replace("", pd.NA).dropna().nunique()) if not final_report.empty else 0
    unsupported_cnt = int((final_report["Supported"] == "No").sum())
    no_fuel_cnt = int(len(no_fuel_df))
    strong_cnt = int((final_report["Best_Status"] == "Strong Match").sum())
    possible_cnt = int((final_report["Best_Status"] == "Possible Match").sum())
    review_needed_cnt = int(len(review_df))

    metrics = {
        "Total vehicles in list": total_input_rows,
        "Total unique vehicle entries": unique_entries,
        "Total brands in list": unique_brands,
        "Total unique models": unique_models,
        "Vehicles supported by one or more lists but no fuel reading": no_fuel_cnt,
        "Total number of vehicles not included in any support list": unsupported_cnt,
        "Support coverage %": round(((unique_entries - unsupported_cnt) / unique_entries) * 100, 1) if unique_entries else 0,
        "Strong match vehicles": strong_cnt,
        "Possible match vehicles": possible_cnt,
        "Vehicles recommended for review": review_needed_cnt,
        "Duplicate input rows": max(0, total_input_rows - unique_entries),
    }

    return metrics, final_report, unsupported_df, no_fuel_df, brand_gap_df, review_df

def write_business_report_workbook(metrics, final_report, unsupported_df, no_fuel_df, brand_gap_df, review_df):
    from openpyxl import Workbook
    wb_rep = Workbook()
    ws_sum = wb_rep.active
    ws_sum.title = "Summary_KPI"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)

    ws_sum["A1"] = "Vehicle Compatibility Clean Report"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A3"] = "KPI"; ws_sum["B3"] = "Value"
    for c in ("A3","B3"):
        ws_sum[c].fill = header_fill; ws_sum[c].font = white_font
    row = 4
    for k, v in metrics.items():
        ws_sum.cell(row,1).value = k
        ws_sum.cell(row,2).value = v
        row += 1

    sections = [
        ("Vehicles not in any support list", unsupported_df.head(25)),
        ("Supported but no fuel reading", no_fuel_df.head(25)),
        ("Brand-wise gaps", brand_gap_df.head(25)),
        ("Recommended review list", review_df.head(25)),
    ]
    for title, df in sections:
        row += 1
        ws_sum.cell(row,1).value = title
        ws_sum.cell(row,1).fill = sub_fill
        ws_sum.cell(row,1).font = bold
        row += 1
        if df.empty:
            ws_sum.cell(row,1).value = "None"
            row += 1
            continue
        for idx, h in enumerate(df.columns, start=1):
            ws_sum.cell(row, idx).value = h
            ws_sum.cell(row, idx).fill = header_fill
            ws_sum.cell(row, idx).font = white_font
        row += 1
        for rec in df.itertuples(index=False):
            for idx, v in enumerate(rec, start=1):
                ws_sum.cell(row, idx).value = v
            row += 1

    def add_df_sheet(name, df):
        ws = wb_rep.create_sheet(name)
        for idx, h in enumerate(df.columns, start=1):
            ws.cell(1, idx).value = h
            ws.cell(1, idx).fill = header_fill
            ws.cell(1, idx).font = white_font
        for r_idx, rec in enumerate(df.itertuples(index=False), start=2):
            for c_idx, v in enumerate(rec, start=1):
                ws.cell(r_idx, c_idx).value = v
        ws.freeze_panes = "A2"
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        autosize_sheet(ws)

    add_df_sheet("Final_Report", final_report)
    add_df_sheet("Unsupported_Vehicles", unsupported_df)
    add_df_sheet("No_Fuel_Vehicles", no_fuel_df)
    add_df_sheet("Brand_Gaps", brand_gap_df)
    add_df_sheet("Needs_Review", review_df)

    for row in ws_sum.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    autosize_sheet(ws_sum)
    wb_rep.save(FINAL_REPORT_XLSX)

def write_report_sheets_to_workbook(wb, metrics, final_report, unsupported_df, no_fuel_df, brand_gap_df, review_df):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)

    ws_sum = prepare_report_sheet(wb, "Summary_KPI", ["KPI","Value"])
    ws_fr = prepare_report_sheet(wb, "Final_Report", list(final_report.columns))
    ws_uv = prepare_report_sheet(wb, "Unsupported_Vehicles", list(unsupported_df.columns))
    ws_nf = prepare_report_sheet(wb, "No_Fuel_Vehicles", list(no_fuel_df.columns))
    ws_bg = prepare_report_sheet(wb, "Brand_Gaps", list(brand_gap_df.columns))
    ws_nr = prepare_report_sheet(wb, "Needs_Review", list(review_df.columns))

    ws_sum["A1"] = "Vehicle Compatibility Clean Report"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A3"] = "KPI"; ws_sum["B3"] = "Value"
    for c in ("A3","B3"):
        ws_sum[c].fill = header_fill
        ws_sum[c].font = white_font
    row = 4
    for k, v in metrics.items():
        ws_sum.cell(row,1).value = k
        ws_sum.cell(row,2).value = v
        row += 1

    sections = [
        ("Vehicles not in any support list", unsupported_df.head(25)),
        ("Supported but no fuel reading", no_fuel_df.head(25)),
        ("Brand-wise gaps", brand_gap_df.head(25)),
        ("Recommended review list", review_df.head(25)),
    ]
    for title, df in sections:
        row += 1
        ws_sum.cell(row,1).value = title
        ws_sum.cell(row,1).fill = sub_fill
        ws_sum.cell(row,1).font = bold
        row += 1
        if df.empty:
            ws_sum.cell(row,1).value = "None"
            row += 1
            continue
        for idx, h in enumerate(df.columns, start=1):
            ws_sum.cell(row, idx).value = h
            ws_sum.cell(row, idx).fill = header_fill
            ws_sum.cell(row, idx).font = white_font
        row += 1
        for rec in df.itertuples(index=False):
            for idx, v in enumerate(rec, start=1):
                ws_sum.cell(row, idx).value = v
            row += 1

    for ws, df in [(ws_fr, final_report), (ws_uv, unsupported_df), (ws_nf, no_fuel_df), (ws_bg, brand_gap_df), (ws_nr, review_df)]:
        for idx, h in enumerate(df.columns, start=1):
            ws.cell(1, idx).value = h
            ws.cell(1, idx).fill = header_fill
            ws.cell(1, idx).font = white_font
        for r_idx, rec in enumerate(df.itertuples(index=False), start=2):
            for c_idx, v in enumerate(rec, start=1):
                ws.cell(r_idx, c_idx).value = v

    for ws in [ws_sum, ws_fr, ws_uv, ws_nf, ws_bg, ws_nr]:
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        autosize_sheet(ws)

def main():
    wb = load_workbook(WORKBOOK)
    ws_in = wb["Vehicle_Input"]
    ws_match = wb["Match_Results"]
    ws_source = wb["Per_Source_Results"]
    ws_top3 = wb["Top3_Candidates"]

    master = normalize_master(load_sheet_df("Supported_Master"))
    alias_df = load_sheet_df("Aliases_Rules")
    overrides = load_sheet_df("Manual_Overrides")
    try:
        gen_rules = load_sheet_df("Generation_Rules")
    except Exception:
        gen_rules = pd.DataFrame()
    try:
        chassis_rules = load_sheet_df("Chassis_Rules")
    except Exception:
        chassis_rules = pd.DataFrame()

    make_aliases, model_aliases, category_aliases = build_alias_maps(alias_df)

    for ws in [ws_match, ws_source, ws_top3]:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for c in row:
                c.value = None

    match_row_ptr = 2
    source_row_ptr = 2
    top_row_ptr = 2
    all_records = []
    input_rows = []

    brand_groups = {brand: g.copy() for brand, g in master.groupby("Brand_Norm", dropna=False)}

    for r in range(2, ws_in.max_row + 1):
        input_id = ws_in[f"A{r}"].value
        raw_desc = ws_in[f"B{r}"].value
        if raw_desc is None or str(raw_desc).strip() == "":
            continue

        raw_desc = str(raw_desc).strip()
        input_rows.append(raw_desc)

        override = check_override(raw_desc, overrides)
        parsed_year = extract_year(raw_desc)
        parsed_gen = extract_generation(raw_desc)
        parsed_make = detect_make(norm_text(raw_desc), master, make_aliases)
        input_category = infer_input_category(raw_desc)
        parsed_family = extract_parsed_model_family(raw_desc, parsed_make, master, model_aliases)
        search_text = build_search_context(raw_desc, parsed_make, input_category, category_aliases)
        if not search_text:
            search_text = parsed_family

        if override:
            rows = master[
                (master["Support_List_Type"].astype(str) == str(override.get("Override_Source",""))) &
                (master["Brand"].astype(str) == str(override.get("Override_Brand",""))) &
                (master["Model"].astype(str) == str(override.get("Override_Model","")))
            ].copy()
        else:
            rows = brand_groups.get(parsed_make, master.iloc[0:0].copy()) if parsed_make else master.iloc[0:0].copy()
            rows = prepare_candidate_pool(rows, input_category, search_text)

        scored = []
        for _, row in rows.iterrows():
            score_info = score_candidate(search_text, parsed_make, parsed_family, parsed_year, parsed_gen, input_category, row)
            if not score_info:
                continue
            status, conf = classify_candidate(score_info, rows, parsed_family, search_text)
            reason_detail = (
                f"Brand locked to '{parsed_make or ''}'. Input category='{input_category or ''}'. "
                f"Search text='{search_text or ''}'. Parsed family='{parsed_family or ''}'. "
                f"Compared against brand='{row.get('Brand','')}', model='{row.get('Model','')}', "
                f"category='{row.get('Vehicle_Category','')}', year range='{row.get('Model_Year_Text','')}'. "
                f"Checks -> brand={score_info['brand_check']}, category={score_info['category_check']}, "
                f"model={score_info['model_check']}, year={score_info['year_check']}; "
                f"text_ratio={score_info['best_ratio']}, token_coverage={score_info['coverage']}."
            )
            rec = {
                "Input_ID": input_id,
                "Raw_Vehicle_Description": raw_desc,
                "Parsed_Brand": parsed_make,
                "Parsed_Model_Family": parsed_family,
                "Parsed_Year": parsed_year if parsed_year is not None else "",
                "Parsed_Generation": parsed_gen,
                "Input_Category": input_category,
                "Source": row["Support_List_Type"],
                "Matched_Brand": row["Brand"],
                "Matched_Model": row["Model"],
                "Matched_Vehicle_Category": row.get("Vehicle_Category",""),
                "Supported_Year_Range": row["Model_Year_Text"],
                "Brand_Check": score_info["brand_check"],
                "Model_Check": score_info["model_check"],
                "Year_Check": score_info["year_check"],
                "Category_Check": score_info["category_check"],
                "Supported_Parameter_Count": row.get("Supported_Parameter_Count",""),
                "Fuel_Data_Available": row.get("Fuel_Data_Available","No"),
                "Parameter_Preview": row.get("Param_Preview",""),
                "Full_Supported_Parameters": row.get("Full_Supported_Parameters",""),
                "Match_Score": score_info["score"],
                "Status": status,
                "Confidence": conf,
                "Reason": overall_reason(score_info["brand_check"], score_info["model_check"], score_info["year_check"], score_info["category_check"]),
                "Reason_Detail": reason_detail,
            }
            scored.append(rec)

        scored = sorted(scored, key=lambda x: (x["Match_Score"], {"High":3,"Medium":2,"Low":1}.get(x["Confidence"],0)), reverse=True)

        # anti-hallucination controls on top ranked candidates
        best_overall = None
        if scored:
            top = scored[0]
            runner = scored[1] if len(scored) > 1 else None
            generic = generic_input_flag(parsed_family, search_text)
            same_family_conflict = runner and norm_text(top["Matched_Model"]) != norm_text(runner["Matched_Model"]) and abs(top["Match_Score"] - runner["Match_Score"]) < 6
            if not parsed_make:
                top["Status"] = "No Reliable Match"; top["Confidence"] = "Low"
            elif generic and top["Status"] == "Possible Match" and same_family_conflict:
                top["Status"] = "Review Needed"; top["Confidence"] = "Low"
            elif top["Match_Score"] < 58 or top["Model_Check"] == "Weak":
                top["Status"] = "No Reliable Match"; top["Confidence"] = "Low"
            best_overall = top

        if not best_overall:
            best_overall = {
                "Parsed_Brand": parsed_make, "Parsed_Model_Family": parsed_family,
                "Parsed_Year": parsed_year or "", "Parsed_Generation": parsed_gen,
                "Input_Category": input_category,
                "Status":"No Reliable Match","Confidence":"Low","Source":"","Matched_Brand":"",
                "Matched_Model":"","Matched_Vehicle_Category":"","Supported_Year_Range":"",
                "Brand_Check":"No" if not parsed_make else "Exact","Model_Check":"Weak","Year_Check":"Not provided" if parsed_year is None else "",
                "Reason":"No candidates found","Reason_Detail":"No candidates found for the given input.",
                "Supported_Parameter_Count":"","Fuel_Data_Available":"No","Parameter_Preview":"","Full_Supported_Parameters":"","Match_Score":0
            }

        ws_match.cell(match_row_ptr, 1).value = input_id
        ws_match.cell(match_row_ptr, 2).value = raw_desc
        ws_match.cell(match_row_ptr, 3).value = best_overall["Parsed_Brand"]
        ws_match.cell(match_row_ptr, 4).value = best_overall["Parsed_Model_Family"]
        ws_match.cell(match_row_ptr, 5).value = best_overall["Parsed_Year"]
        ws_match.cell(match_row_ptr, 6).value = best_overall["Parsed_Generation"]
        ws_match.cell(match_row_ptr, 7).value = best_overall["Status"]
        ws_match.cell(match_row_ptr, 8).value = best_overall["Confidence"]
        ws_match.cell(match_row_ptr, 9).value = best_overall["Source"]
        ws_match.cell(match_row_ptr,10).value = best_overall["Matched_Brand"]
        ws_match.cell(match_row_ptr,11).value = best_overall["Matched_Model"]
        ws_match.cell(match_row_ptr,12).value = best_overall["Supported_Year_Range"]
        ws_match.cell(match_row_ptr,13).value = best_overall["Brand_Check"]
        ws_match.cell(match_row_ptr,14).value = best_overall["Model_Check"]
        ws_match.cell(match_row_ptr,15).value = best_overall["Year_Check"]
        ws_match.cell(match_row_ptr,16).value = best_overall["Reason"]
        ws_match.cell(match_row_ptr,17).value = action_hint(best_overall["Status"], parsed_year)
        ws_match.cell(match_row_ptr,18).value = best_overall["Supported_Parameter_Count"]
        ws_match.cell(match_row_ptr,19).value = best_overall["Fuel_Data_Available"]
        ws_match.cell(match_row_ptr,20).value = best_overall["Parameter_Preview"]
        match_row_ptr += 1

        for source in sorted({x["Source"] for x in scored if x.get("Source")}):
            candidates = [x for x in scored if x["Source"] == source]
            if not candidates:
                continue
            best = candidates[0]
            vals = [
                input_id, raw_desc, source, best["Status"], best["Confidence"],
                best["Parsed_Brand"], best["Parsed_Model_Family"], best["Parsed_Year"],
                best["Matched_Brand"], best["Matched_Model"], best["Supported_Year_Range"],
                best["Brand_Check"], best["Model_Check"], best["Year_Check"],
                best["Supported_Parameter_Count"], best["Fuel_Data_Available"], best["Parameter_Preview"], best["Reason"],
                best["Reason_Detail"], best["Full_Supported_Parameters"]
            ]
            for cidx, v in enumerate(vals, start=1):
                ws_source.cell(source_row_ptr, cidx).value = v
            source_row_ptr += 1

        for rank, cand in enumerate(scored[:3], start=1):
            vals = [
                input_id, raw_desc, rank, cand["Source"], cand["Status"], cand["Confidence"], cand["Match_Score"],
                cand["Matched_Brand"], cand["Matched_Model"], cand["Supported_Year_Range"],
                cand["Brand_Check"], cand["Model_Check"], cand["Year_Check"], cand["Reason"]
            ]
            for cidx, v in enumerate(vals, start=1):
                ws_top3.cell(top_row_ptr, cidx).value = v
            top_row_ptr += 1

        if scored:
            all_records.extend(scored)
        else:
            all_records.append({
                "Input_ID": input_id, "Raw_Vehicle_Description": raw_desc, "Parsed_Brand": parsed_make,
                "Parsed_Model_Family": parsed_family, "Parsed_Year": parsed_year if parsed_year is not None else "",
                "Parsed_Generation": parsed_gen, "Input_Category": input_category,
                "Source": "", "Matched_Brand": "", "Matched_Model": "", "Matched_Vehicle_Category": "",
                "Supported_Year_Range": "", "Brand_Check": "No" if not parsed_make else "Exact",
                "Model_Check": "Weak", "Year_Check": "Not provided" if parsed_year is None else "",
                "Category_Check": "Not provided", "Supported_Parameter_Count": "", "Fuel_Data_Available": "No",
                "Parameter_Preview": "", "Full_Supported_Parameters": "", "Match_Score": 0,
                "Status": best_overall["Status"], "Confidence": best_overall["Confidence"],
                "Reason": best_overall["Reason"], "Reason_Detail": best_overall["Reason_Detail"]
            })

        supported = "Yes" if best_overall["Status"] in ("Strong Match", "Possible Match") else "No"
        ws_in[f"D{r}"] = "Yes"
        ws_in[f"E{r}"] = f"Processed - {supported}"

    metrics, vehicle_summary, adaptor_detail, adaptor_summary, brand_summary, unsupported_df = build_business_report(all_records, input_rows)

    color_results(ws_match, "G")
    color_results(ws_source, "D")
    color_results(ws_top3, "E")

    update_table_ref(ws_match, "T_MatchResults")
    update_table_ref(ws_source, "T_PerSourceResults")
    update_table_ref(ws_top3, "T_Top3Candidates")

    write_report_sheets_to_workbook(wb, metrics, vehicle_summary, adaptor_detail, adaptor_summary, brand_summary, unsupported_df)
    wb.save(WORKBOOK)
    write_business_report_workbook(metrics, vehicle_summary, adaptor_detail, adaptor_summary, brand_summary, unsupported_df)
    if not vehicle_summary.empty:
        vehicle_summary.to_csv(FINAL_REPORT_CSV, index=False)

    print(f"Processed {len(input_rows)} input row(s).")
    print(f"Per-source rows written: {max(0, source_row_ptr-2)}")
    print(f"Top-3 candidate rows written: {max(0, top_row_ptr-2)}")
    print(f"Smart report saved to: {FINAL_REPORT_XLSX.resolve()}")
    print(f"Workbook saved to: {WORKBOOK.resolve()}")

if __name__ == "__main__":
    main()
